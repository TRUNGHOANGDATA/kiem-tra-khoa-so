"""Màn "Thay đổi từ khi chốt" — chứng từ (thêm/bớt/SỬA) + chỉ số CĐPS, so với bản chốt.

Khác `lay_diff_chot` (chỉ liệt kê dòng thêm/bớt rời rạc của bảng kê): ở đây gom theo
chứng từ và kèm cả CĐPS, để kiểm soát bút toán phát sinh sau khi đã khóa sổ.
"""
import pandas as pd

from tests.test_nhieu_chi_nhanh import _api_hai_chi_nhanh


def _api(tmp_path):
    """Hai chi nhánh; `_api_hai_chi_nhanh` đã ép kho vào thư mục test."""
    return _api_hai_chi_nhanh(tmp_path)


def _chot_het(api):
    for i in range(len(api._dv)):
        api.chon_don_vi(i)
        api.chot_so()


def test_chua_chot_thi_bao_ro(tmp_path):
    api = _api(tmp_path)
    api.chay_kiem_tra()
    assert "chưa chốt" in api.thay_doi_tu_khi_chot()["loi"].lower()


def test_chot_xong_khong_doi_gi_thi_sach(tmp_path):
    api = _api(tmp_path)
    api.chay_kiem_tra()
    _chot_het(api)
    r = api.thay_doi_tu_khi_chot()
    assert "loi" not in r
    assert r["don_vi"][0]["chung_tu"]["tom_tat"] == {"ct_them": 0, "ct_bot": 0, "ct_sua": 0}


def test_sua_mot_dong_hien_thanh_chung_tu_BI_SUA(tmp_path):
    """Đổi số tiền một dòng sau khi chốt: phải ra 1 chứng từ SỬA, kèm dòng cũ và mới."""
    api = _api(tmp_path)
    api.chay_kiem_tra()
    _chot_het(api)
    api.chon_don_vi(0)
    d = api._hien
    d.df = d.df.copy()
    d.df.loc[d.df.index[0], "Amount"] = float(d.df.iloc[0]["Amount"]) + 1_000

    r = api.thay_doi_tu_khi_chot()
    ct = r["don_vi"][0]["chung_tu"]
    assert ct["tom_tat"] == {"ct_them": 0, "ct_bot": 0, "ct_sua": 1}
    sua = ct["sua"][0]
    assert len(sua["dong_cu"]) == 1 and len(sua["dong_moi"]) == 1
    assert sua["dong_moi"][0]["Amount"] - sua["dong_cu"][0]["Amount"] == 1_000


def test_them_chung_tu_moi_thi_dem_vao_THEM(tmp_path):
    api = _api(tmp_path)
    api.chay_kiem_tra()
    _chot_het(api)
    api.chon_don_vi(0)
    d = api._hien
    moi = d.df.iloc[[0]].copy()
    moi["DocNo"] = "CT-MOI-999"
    d.df = pd.concat([d.df, moi], ignore_index=True)

    ct = api.thay_doi_tu_khi_chot()["don_vi"][0]["chung_tu"]
    assert ct["tom_tat"]["ct_them"] == 1 and ct["tom_tat"]["ct_sua"] == 0
    assert ct["them"][0]["DocNo"] == "CT-MOI-999"


def test_pham_vi_tat_ca_gom_moi_chi_nhanh(tmp_path):
    api = _api(tmp_path)
    api.chay_kiem_tra()
    _chot_het(api)
    r = api.thay_doi_tu_khi_chot("tat_ca")
    assert len(r["don_vi"]) == 2
    assert {x["chi_nhanh"] for x in r["don_vi"]} == {"A01", "B02"}
    assert r["tom_tat"]["so_chi_nhanh"] == 2


def test_ban_chot_khong_kem_cdps_thi_noi_ro_chua_co_moc(tmp_path):
    """Chốt khi chưa nạp CĐPS -> không có mốc so. Phải báo `co_cdps_chot=False`, KHÔNG
    được im lặng coi như "chỉ số không đổi" (luật: không kết luận từ thứ không có)."""
    api = _api(tmp_path)
    api.chay_kiem_tra()
    _chot_het(api)
    dv = api.thay_doi_tu_khi_chot("tat_ca")["don_vi"]
    assert all(x["co_cdps_chot"] is False and x["cdps"] is None for x in dv)


def test_xuat_excel_hai_sheet_ghep_ten_chi_nhanh(tmp_path):
    """Bằng chứng kiểm soát để gửi đi: mỗi dòng tự mang tên chi nhánh."""
    import openpyxl

    from app import report
    api = _api(tmp_path)
    api.chay_kiem_tra()
    _chot_het(api)
    api.chon_don_vi(0)
    d = api._hien
    d.df = d.df.copy()
    d.df.loc[d.df.index[0], "Amount"] = float(d.df.iloc[0]["Amount"]) + 1_000

    p = report.xuat_thay_doi(api.thay_doi_tu_khi_chot("tat_ca"), str(tmp_path / "out"))
    wb = openpyxl.load_workbook(p)
    assert wb.sheetnames == ["Chung tu thay doi", "CDPS thay doi"]
    ws = wb["Chung tu thay doi"]
    tieu_de = [c.value for c in ws[4]]
    assert tieu_de[0] == "Chi nhánh"
    nhan = {ws.cell(r, tieu_de.index("Thay đổi") + 1).value for r in range(5, ws.max_row + 1)}
    assert nhan == {"✎ Sửa — trước", "✎ Sửa — sau"}     # có ký hiệu, không chỉ dựa vào màu


def test_xuat_excel_noi_ro_khi_ban_chot_khong_kem_cdps(tmp_path):
    import openpyxl

    from app import report
    api = _api(tmp_path)
    api.chay_kiem_tra()
    _chot_het(api)
    p = report.xuat_thay_doi(api.thay_doi_tu_khi_chot("tat_ca"), str(tmp_path / "out"))
    ws = openpyxl.load_workbook(p)["CDPS thay doi"]
    ly_do = [ws.cell(r, 1).value for r in range(5, ws.max_row + 1)]
    assert len(ly_do) == 2                              # hai chi nhánh, mỗi cái một dòng
    assert "Chưa có mốc" in str([c.value for c in ws[5]])
